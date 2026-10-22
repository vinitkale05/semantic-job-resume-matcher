"""Extract text from PDF and DOCX resumes."""
from io import BytesIO
from pathlib import Path
try:
    import pymupdf as fitz
except ImportError:
    import fitz
from docx import Document


def parse_resume(uploaded_file) -> str:
    """Parse a Streamlit UploadedFile, or a filesystem path, into text."""
    if hasattr(uploaded_file, "getvalue"):
        data = uploaded_file.getvalue()
        name = getattr(uploaded_file, "name", "resume.pdf")
    elif isinstance(uploaded_file, (str, Path)):
        path = Path(uploaded_file)
        data = path.read_bytes()
        name = path.name
    else:
        raise TypeError("Expected an uploaded file or a filesystem path.")

    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        with fitz.open(stream=data, filetype="pdf") as pdf:
            text = "\n".join(page.get_text("text") for page in pdf)
    elif suffix == ".docx":
        document = Document(BytesIO(data))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(paragraphs)
    else:
        raise ValueError("Unsupported file type. Use PDF or DOCX.")

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in the document.")
    return text
